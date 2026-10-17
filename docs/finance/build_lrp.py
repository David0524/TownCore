"""TownCore long-range plan (FY2027-FY2031): a linked, formula-driven workbook.

Every number on the calculation tabs is a formula over the Assumptions tab, so changing an input
(or the scenario switch) recalculates the whole plan.
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'TownCore-LRP.xlsx')
YEARS = ['2027', '2028', '2029', '2030', '2031']
AGES = ['Age 1', 'Age 2', 'Age 3', 'Age 4', 'Age 5']
COLS = ['C', 'D', 'E', 'F', 'G']  # five periods (years or ages) always live in C..G

ARIAL = 'Arial'
F_BASE = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_INPUT = Font(name=ARIAL, size=10, color='0000FF')
F_LINK = Font(name=ARIAL, size=10, color='008000')
F_TITLE = Font(name=ARIAL, size=14, bold=True)
F_SUB = Font(name=ARIAL, size=10, italic=True, color='555555')
F_HEAD = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
F_NOTE = Font(name=ARIAL, size=9, color='555555')
FILL_HEAD = PatternFill('solid', start_color='16233A')
FILL_KEY = PatternFill('solid', start_color='FFFF00')
FILL_SECTION = PatternFill('solid', start_color='E8ECF2')
THIN = Side(style='thin', color='999999')
TOP = Border(top=THIN)

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x"'

wb = Workbook()
REF = {}  # name -> absolute reference usable from any sheet


def q(sheet):
    return f"'{sheet}'" if ' ' in sheet else sheet


def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def header(ws, row, label, periods):
    ws.cell(row=row, column=1, value=label).font = F_HEAD
    for c in range(1, 9):
        ws.cell(row=row, column=c).fill = FILL_HEAD
    for i, p in enumerate(periods):
        cell = ws[f'{COLS[i]}{row}']
        cell.value = p
        cell.font = F_HEAD
        cell.alignment = Alignment(horizontal='right')


def section(ws, row, label):
    ws.cell(row=row, column=1, value=label).font = F_BOLD
    for c in range(1, 9):
        ws.cell(row=row, column=c).fill = FILL_SECTION


def title(ws, text, sub):
    ws['A1'] = text
    ws['A1'].font = F_TITLE
    ws['A2'] = sub
    ws['A2'].font = F_SUB


# ----------------------------------------------------------------------------------------------
# Assumptions
A = wb.active
A.title = 'Assumptions'
title(A, 'TownCore LRP: assumptions', 'Blue = input you can change. Yellow = key assumption. '
      'Every other tab is formulas over this one. Columns C-G are years (2027-2031) or a town\'s '
      'age in years (Age 1 = launch year), as each block\'s header says.')
setw(A, {'A': 46, 'B': 12, 'C': 12, 'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 70})
row = 4


def single(name, label, value, fmt, note, key=False, unit=''):
    """One input in column C (B holds the unit)."""
    global row
    A.cell(row=row, column=1, value=label).font = F_BASE
    A.cell(row=row, column=2, value=unit).font = F_NOTE
    c = A[f'C{row}']
    c.value = value
    c.font = F_INPUT
    c.number_format = fmt
    if key:
        c.fill = FILL_KEY
    A[f'H{row}'] = note
    A[f'H{row}'].font = F_NOTE
    REF[name] = f'Assumptions!$C${row}'
    row += 1


def series(name, label, values, fmt, note, key=False, unit=''):
    """Five inputs across C..G (years or ages)."""
    global row
    A.cell(row=row, column=1, value=label).font = F_BASE
    A.cell(row=row, column=2, value=unit).font = F_NOTE
    for i, v in enumerate(values):
        c = A[f'{COLS[i]}{row}']
        c.value = v
        c.font = F_INPUT
        c.number_format = fmt
        if key:
            c.fill = FILL_KEY
    A[f'H{row}'] = note
    A[f'H{row}'].font = F_NOTE
    REF[name] = row  # row number; columns C..G
    row += 1


def blank():
    global row
    row += 1


# Scenario switch
section(A, row, 'Scenario')
row += 1
A.cell(row=row, column=1, value='Active scenario (1 = Downside, 2 = Base, 3 = Upside)').font = F_BOLD
A[f'C{row}'] = 2
A[f'C{row}'].font = F_INPUT
A[f'C{row}'].fill = FILL_KEY
A[f'H{row}'] = 'Pick 1, 2 or 3. The multipliers below scale launches, audience, monetization and unit costs.'
A[f'H{row}'].font = F_NOTE
REF['scenario'] = f'Assumptions!$C${row}'
row += 1
for c, h in zip(['C', 'D', 'E', 'F'], ['Downside', 'Base', 'Upside', 'Active']):
    A[f'{c}{row}'] = h
    A[f'{c}{row}'].font = F_BOLD
    A[f'{c}{row}'].alignment = Alignment(horizontal='right')
row += 1
for name, label, vals, note in [
    ('m_launch', 'Town launches multiplier', [0.6, 1.0, 1.3], 'Scales the launch schedule.'),
    ('m_aud', 'Audience multiplier', [0.7, 1.0, 1.25], 'Scales subscriber penetration and pageviews.'),
    ('m_mon', 'Monetization multiplier', [0.7, 1.0, 1.2], 'Scales fill, sell-through and conversion rates.'),
    ('m_cost', 'Unit cost multiplier', [1.2, 1.0, 0.9], 'Scales per-town direct costs and launch costs.'),
]:
    A.cell(row=row, column=1, value=label).font = F_BASE
    for c, v in zip(['C', 'D', 'E'], vals):
        A[f'{c}{row}'] = v
        A[f'{c}{row}'].font = F_INPUT
        A[f'{c}{row}'].number_format = MULT
    A[f'F{row}'] = f'=INDEX(C{row}:E{row},1,{REF["scenario"]})'
    A[f'F{row}'].number_format = MULT
    A[f'F{row}'].font = F_BOLD
    A[f'H{row}'] = note
    A[f'H{row}'].font = F_NOTE
    REF[name] = f'Assumptions!$F${row}'
    row += 1
blank()

# Rollout
section(A, row, 'Launch capacity: how many towns can launch, and how fast')
row += 1
single('addressable', 'US towns with more than 2,000 people', 8433, NUM,
       'Owner plan (2026-10-17): every US town over 2,000 people. Census Vintage 2025 '
       '(sub-est2025.csv): 7,683 incorporated places (SUMLEV 162; Florida 295) plus 750 New '
       'England towns (active county subdivisions with no place of their own).', unit='towns')
single('ready_share', 'Share that pass readiness (Rule 7)', 1.0, PCT,
       'Owner, 2026-10-17: every town in America can make 3 real items a day; Rule 7 is the target '
       'we build sources toward. Today Boca makes about 0.8 a day and Delray 0.53, so the source work '
       'continues. Set lower to see the plan if some towns stay short.',
       key=True)
single('name_share', 'Share allowed by the naming rule', 1.0, PCT,
       'Owner, 2026-10-17: naming never stops a launch. 22 patterns; a neighbor\'s pattern is '
       'used only when nothing else is left.')
row += 1
header(A, row, 'Fiscal year', YEARS)
row += 1
series('launch_cap', 'Launches the team can do in a year', [3000, 3000, 3000, 3000, 3000], NUM,
       'About 250 a month: the launcher is automated, each launch needs ~45 minutes of people '
       '(approval, Facebook page, Google Publisher Center). A pacing choice too: launching '
       'hundreds of AI-written sites at once invites Google\'s scaled-content review. Scaled by '
       'the launch multiplier.', key=True, unit='towns')
single('per_launcher', 'Launches per launch operator per year', 400, NUM,
       '~45 minutes of hands-on work per launch, plus fixes and follow-up. Estimate.', unit='towns')
blank()
section(A, row, 'Town size mix and size-driven inputs (CLAUDE.md §5 tiers)')
row += 1
for c, h in zip(['C', 'D', 'E', 'F'], ['Small', 'Mid', 'Large', 'Weighted']):
    A[f'{c}{row}'] = h
    A[f'{c}{row}'].font = F_BOLD
    A[f'{c}{row}'].alignment = Alignment(horizontal='right')
row += 1


def tier(name, label, vals, fmt, note, key=False, weighted=True):
    global row
    A.cell(row=row, column=1, value=label).font = F_BASE
    for c, v in zip(['C', 'D', 'E'], vals):
        A[f'{c}{row}'] = v
        A[f'{c}{row}'].font = F_INPUT
        A[f'{c}{row}'].number_format = fmt
        if key:
            A[f'{c}{row}'].fill = FILL_KEY
    if weighted:
        A[f'F{row}'] = f'=SUMPRODUCT(C{row}:E{row},$C${REF["mix"]}:$E${REF["mix"]})'
        A[f'F{row}'].number_format = fmt
        A[f'F{row}'].font = F_BOLD
        REF[name] = f'Assumptions!$F${row}'
    else:
        REF[name] = row
    A[f'H{row}'] = note
    A[f'H{row}'].font = F_NOTE
    row += 1


tier('mix', 'Share of launches', [0.84, 0.12, 0.04], PCT,
     'Census 2025, towns over 2,000: 7,067 under 30,000 (small, incl. the 750 New England towns), '
     '1,023 of 30,000-100,000 (mid), 343 of 100,000+ (large); tiers from packages/config/src/launch.ts.',
     key=True, weighted=False)
A[f'F{REF["mix"]}'] = f'=SUM(C{REF["mix"]}:E{REF["mix"]})'
A[f'F{REF["mix"]}'].number_format = PCT
A[f'F{REF["mix"]}'].font = F_BOLD
tier('households', 'Households per town', [3300, 21000, 70000], NUM,
     'Mean town population by tier (8,397 / 53,365 / 185,762, Census 2025, each town capped at '
     '300,000) / ~2.5 people a household. Revenue scales with households, so the mean is used.')
tier('listings', 'Directory listings per town', [180, 1100, 3500], NUM,
     'Boca live data 2026-10-16: 2,334 active listings (1,979 state licenses + 355 mapped). '
     'Small and mid scaled down; estimate.')
tier('stories', 'Stories per day (story budget)', [3, 6, 10], NUM,
     'CLAUDE.md D-M3-7: daily story budget 3 / 6 / 10 by tier.')
tier('banner', 'Town sponsor banner price ($/month)', [150, 400, 900], USD,
     'Direct-sold home.top sponsor (§11b). Estimate; local banner rates vary widely.')
blank()

# Audience by age
section(A, row, 'Audience (by town age; Age 1 = launch year)')
row += 1
header(A, row, 'Town age', AGES)
row += 1
series('sub_pen', 'Newsletter subscribers as % of households', [0.015, 0.04, 0.06, 0.075, 0.085], PCT,
       'Estimate. Mature local newsletters report reaching roughly 5-15% of households; '
       'scaled by the audience multiplier.', key=True)
series('pv_hh', 'Monthly pageviews per household', [1.0, 2.0, 3.0, 3.5, 4.0], NUM1,
       'Estimate; driven by search, Facebook and the daily newsletter. Scaled by audience multiplier.')
series('survival', 'Share of launched towns still live', [1.0, 0.97, 0.95, 0.94, 0.93], PCT,
       'Rule 7: towns that fall below readiness are paused. Estimate.')
blank()

# Monetization by age
section(A, row, 'Monetization (by town age)')
row += 1
header(A, row, 'Town age', AGES)
row += 1
series('feat_fill', 'Featured directory spots filled (% of 3 per category)',
       [0.05, 0.15, 0.25, 0.32, 0.38], PCT,
       'Estimate. Spots #1/#2/#3 per category (§11). Scaled by monetization multiplier.', key=True)
series('claim', 'Listings claimed (cumulative %)', [0.02, 0.05, 0.08, 0.10, 0.12], PCT,
       'Estimate; claim flow by email confirmation (M6).')
series('enh_conv', 'Claimed listings buying Enhanced (%)', [0.08, 0.12, 0.15, 0.17, 0.18], PCT,
       'Estimate. Scaled by monetization multiplier.', key=True)
series('nl_sell', 'Newsletter sponsorship sell-through', [0.10, 0.30, 0.45, 0.55, 0.60], PCT,
       'Share of issues with a paid presenting sponsor (§11b newsletter.top). Estimate.')
series('ban_sell', 'Town sponsor banner sell-through', [0.10, 0.30, 0.50, 0.60, 0.65], PCT,
       'Share of months the home.top banner is direct-sold. Estimate.')
series('cat_sell', 'Category sponsorships sold (% of categories)', [0.0, 0.02, 0.05, 0.08, 0.10], PCT,
       'category.sponsor / exclusive (§11b); built but off at launch (D-M6-6), so Age 1 = 0.')
series('prog_on', 'Programmatic ads switched on (1 = yes)', [0, 0, 1, 1, 1], NUM,
       'Gated: only after a town passes Gate 2 with steady traffic (§11b). Assumed from Age 3.')
blank()

# Prices
section(A, row, 'Prices and rates')
row += 1
single('categories', 'Directory categories per town', 28, NUM,
       'packages/config/src/directory.ts lists 28 categories.', unit='count')
single('feat_price', 'Featured spots #1 + #2 + #3 per category', 38, USD,
       'Town config pricing.featured_usd [19, 12, 7] (owner, 2026-10-16).', unit='$/month')
single('trial', 'Free trial on a first featured spot', 30, NUM,
       'Town config free_trial_days: 30.', unit='days')
single('trial_loss', 'Revenue lost to trials (share of featured revenue)', 0.08, PCT,
       'Estimate: one free month per new spot over an average spot life of about a year.')
single('enh_price', 'Enhanced listing price', 39, USD, 'Town config pricing.enhanced_usd: 39.',
       unit='$/month')
single('cat_price', 'Category sponsorship price', 99, USD,
       'Estimate; price not yet set by the owner (M6 open item).', unit='$/month')
single('nl_sends', 'Newsletter issues per year', 300, NUM,
       'Daily, but thin days are skipped (D-M5-1); estimate.', unit='issues')
single('nl_cpm', 'Newsletter sponsorship CPM', 30, USD,
       'Estimate for local newsletter sponsorships; verify with first sales.', unit='$/1,000 subs')
single('prog_rpm', 'Programmatic revenue per 1,000 pageviews', 6, USD2,
       'Estimate for a local news site after Gate 2.', unit='$ RPM')
single('prog_fill', 'Pageviews with programmatic ads', 0.8, PCT,
       'Share of pageviews where the rail and inline slots are not direct-sold. Estimate.')
blank()

# Direct costs
section(A, row, 'Direct costs per town')
row += 1
single('ai_story', 'AI cost per published story', 0.15, USD2,
       'Writer, separate fact-check and an AI review pass (owner: article review becomes an AI agent). '
       'Measured ~$0.10 for writer + fact-check on repo pricing; verify against current Anthropic pricing '
       '(CLAUDE.md §8). Scaled by unit cost multiplier.', key=True, unit='$/story')
single('ai_overhead', 'AI overhead per town per day', 0.50, USD2,
       'Classification, revisions, guides, entity extraction. Estimate.', unit='$/day')
single('ai_cap', 'AI spend cap per town per day', 5.0, USD2,
       'CLAUDE.md D-M3-7: $5 daily cap per town, enforced in code.', unit='$/day')
single('hosting', 'Hosting, database and data per town', 25, USD,
       'Vercel, Railway Postgres share, storage, maps. Estimate.', unit='$/month')
single('ses', 'Email sending (Amazon SES)', 0.10, USD2,
       'AWS SES list price per 1,000 emails; verify at aws.amazon.com/ses/pricing.', unit='$/1,000')
single('stripe_pct', 'Card processing, % of directory revenue', 0.029, PCT,
       'Stripe standard US card rate 2.9% + $0.30; verify at stripe.com/pricing.')
single('stripe_fixed', 'Card processing, fixed per charge', 0.30, USD2, 'Stripe standard US.',
       unit='$/charge')
single('domain', 'Domain renewal per town', 20, USD, 'Estimate.', unit='$/year')
single('launch_ai', 'AI cost per launch', 25, USD, 'CLAUDE.md D-M7-11: $25 of AI per launch.',
       unit='$/launch')
single('launch_domain', 'Domain purchase per launch', 15, USD, 'Estimate.', unit='$/launch')
single('launch_mkt', 'Launch audience seeding (paid social)', 0, USD,
       'Owner, 2026-10-17: bootstrapped, no paid launch seeding (was $1,500 a launch). Audience '
       'comes from search, the town Facebook page and the newsletter.', key=True, unit='$/launch')
single('sub_cac', 'Paid cost per new subscriber', 2.0, USD2, 'Estimate.', unit='$/subscriber')
single('sub_paid', 'Share of new subscribers acquired with paid ads', 0.0, PCT,
       'Bootstrapped: no paid acquisition. Subscribers arrive from search, Facebook shares, the '
       'sites\' sign-up prompts and word of mouth.')
blank()

# People
section(A, row, 'People (fully loaded = salary x (1 + benefits load))')
row += 1
single('benefits', 'Benefits and payroll tax load', 0.25, PCT, 'Estimate.')
single('raise', 'Annual salary increase', 0.03, PCT, 'Estimate.')
single('commission', 'Sales commission on direct-sold revenue', 0.0, PCT,
       'Owner: advertisers sign up and pay on the site; no commissioned sales force.')
single('per_reviewer', 'Towns per editorial standards editor', 1000, NUM,
       'Owner, 2026-10-17: article review becomes an AI agent once we trust it. A small team audits '
       'the agent, handles corrections and reader reports. Until the agent ships, Tier 3 still needs '
       'a person (Rule 5).', key=True, unit='towns')
single('per_ae', 'Towns per advertiser partnerships manager', 1000, NUM,
       'Owner, 2026-10-17: sales are self-serve on the site. A small team handles regional and '
       'multi-town advertisers.', key=True, unit='towns')
single('per_support', 'Towns per support and success rep', 120, NUM, 'Estimate.', unit='towns')
single('founder_fte', 'Founder time for operations', 2, NUM1,
       'Owner, 2026-10-17: founders take no salary, at any point. Full-time equivalents of founder '
       'time for support, review and launches; it covers the first needed FTEs at $0. Input: confirm.',
       key=True, unit='FTE')
single('people_share', 'Paid staff budget, share of last month\'s revenue', 0.40, PCT,
       'Bootstrapped hiring rule: staff are hired only out of revenue. Needed work beyond founder '
       'time is paid up to this share of the previous month\'s revenue, and never more than last '
       'month\'s profit before staff; the rest shows as unfunded FTE, and no town launches while '
       'any is unfunded.', key=True)
blank()
section(A, row, 'Base salaries (2027 $)')
row += 1
single('sal_lead', 'Leadership (hires beyond the founders)', 160000, USD,
       'Estimate. Founders are unpaid and not counted here.', unit='$/year')
single('sal_eng', 'Engineering', 170000, USD, 'Estimate.', unit='$/year')
single('sal_rev', 'Editorial standards and launch operations', 70000, USD, 'Estimate.', unit='$/year')
single('sal_ae', 'Advertiser partnerships', 65000, USD, 'Estimate.', unit='$/year')
single('sal_sup', 'Support and success', 55000, USD, 'Estimate.', unit='$/year')
single('sal_ga', 'Finance, legal and operations', 110000, USD, 'Estimate.', unit='$/year')
blank()
section(A, row, 'Central hires wanted (by fiscal year; paid only when the staff budget allows)')
row += 1
header(A, row, 'Fiscal year', YEARS)
row += 1
series('hc_lead', 'Leadership (not founders)', [0, 0, 0, 1, 2], NUM, 'Plan; founders lead unpaid.',
       unit='people')
series('hc_eng', 'Engineering and product', [0, 1, 2, 4, 6], NUM,
       'Plan; founders and Claude Code build until revenue pays for engineers.', unit='people')
series('hc_ga', 'Finance, legal and operations', [0, 0, 1, 1, 2], NUM, 'Plan.', unit='people')
blank()
section(A, row, 'Other operating costs (by fiscal year)')
row += 1
header(A, row, 'Fiscal year', YEARS)
row += 1
series('legal_fixed', 'Legal and media liability insurance (fixed)', [1500, 3000, 6000, 12000, 20000],
       USD, 'Media liability matters for a publisher (§2). Starts with a small policy and grows with the '
       'network. Estimate: get a quote.', key=True, unit='$')
single('brand_share', 'Corporate marketing, share of revenue', 0.02, PCT,
       'Bootstrapped: advertiser marketing paid out of revenue, not ahead of it. Estimate.')
single('platform_base', 'Fixed platform costs', 100, USD,
       'Vercel Pro, Railway base and other subscriptions before any town (per-town hosting is separate). '
       'Estimate: check against the bills.',
       unit='$/month')
single('legal_town', 'Legal and insurance per live town', 100, USD, 'Estimate.', unit='$/town/year')
single('tools', 'Software and tools per employee', 3000, USD, 'Estimate.', unit='$/year')
single('laptop', 'Equipment per new hire (capex)', 2500, USD, 'Estimate.', unit='$/hire')
single('dep_years', 'Equipment depreciation life', 3, NUM, 'Straight line.', unit='years')
blank()
section(A, row, 'Finance')
row += 1
single('tax', 'Income tax rate', 0.25, PCT, 'Federal 21% plus state, simplified; losses carried forward.')
single('nwc', 'Net working capital, % of revenue', 0.03, PCT,
       'Stripe settles in days; mostly prepaid monthly plans. Estimate.')
single('open_cash', 'Opening cash, start of FY2027', 5000, USD,
       'Owner, 2026-10-17: about $5k in cash; bootstrapped, no outside money.', key=True, unit='$')
single('cash_floor', 'Minimum cash kept in the bank', 1000, USD,
       'Launches stop before cash would fall below this. Input.', key=True, unit='$')
single('buffer_months', 'Months of fixed costs and staff kept in the bank', 4, NUM,
       'Launches stop unless cash covers the floor, this many months of fixed costs and paid staff, '
       'and every young town\'s remaining burn until it pays its own way (Monthly sheet).',
       unit='months')
row += 1
header(A, row, 'Fiscal year', YEARS)
row += 1
series('equity', 'Equity raised', [0, 0, 0, 0, 0], USD,
       'Owner, 2026-10-17: bootstrapped, no outside money. Added in the first month of its year.',
       key=True, unit='$')
LAST_A = row

# ----------------------------------------------------------------------------------------------
# Rollout
R = wb.create_sheet('Rollout')  # filled after the Monthly sheet, which drives it

# ----------------------------------------------------------------------------------------------
# Unit economics (one average town, by age)
U = wb.create_sheet('Unit Economics')
title(U, 'Unit economics: one average launched town, by age',
      'Age 1 is the launch year (half a year live on average). Values per launched town, already '
      'weighted by the share of towns still live. Network totals on Revenue and Costs multiply '
      'these by each year\'s launch cohorts.')
setw(U, {'A': 50, 'B': 4, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 60})
header(U, 4, 'Town age', AGES)
urow = 5
UREF = {}


def urow_add(name, label, formula_for_col, fmt, note='', bold=False, top=False):
    """formula_for_col(col_letter, index) -> formula string."""
    global urow
    U.cell(row=urow, column=1, value=label).font = F_BOLD if bold else F_BASE
    for i, c in enumerate(COLS):
        cell = U[f'{c}{urow}']
        cell.value = formula_for_col(c, i)
        cell.number_format = fmt
        if bold:
            cell.font = F_BOLD
        if top:
            cell.border = TOP
    if note:
        U[f'H{urow}'] = note
        U[f'H{urow}'].font = F_NOTE
    UREF[name] = urow
    urow += 1


def usec(label):
    global urow
    section(U, urow, label)
    urow += 1


def a(name, c):
    return f'Assumptions!{c}${REF[name]}'


usec('Time live and survival')
urow_add('live_frac', 'Share of the year live', lambda c, i: '=0.5' if i == 0 else '=1', PCT,
         'Launches spread across the year, so a town is live about half its launch year.')
urow_add('surv', 'Share of launched towns still live', lambda c, i: f'={a("survival", c)}', PCT)
urow_add('town_years', 'Live town-years per launched town',
         lambda c, i: f'={c}{UREF["live_frac"]}*{c}{UREF["surv"]}', NUM1 if False else '0.00')
usec('Audience')
urow_add('subs', 'Newsletter subscribers (year end)',
         lambda c, i: f'={REF["households"]}*{a("sub_pen", c)}*{REF["m_aud"]}*{c}{UREF["surv"]}', NUM)
urow_add('new_subs', 'New subscribers in the year',
         lambda c, i: f'={c}{UREF["subs"]}' if i == 0 else f'=MAX(0,{c}{UREF["subs"]}-{COLS[i-1]}{UREF["subs"]})',
         NUM)
urow_add('pv', 'Pageviews in the year',
         lambda c, i: f'={REF["households"]}*{a("pv_hh", c)}*{REF["m_aud"]}*12*{c}{UREF["town_years"]}', NUM)
usec('Directory')
urow_add('feat_spots', 'Featured spots sold (average)',
         lambda c, i: f'={REF["categories"]}*3*MIN(1,{a("feat_fill", c)}*{REF["m_mon"]})*{c}{UREF["surv"]}',
         NUM1)
urow_add('enh_biz', 'Enhanced listings (average)',
         lambda c, i: f'={REF["listings"]}*{a("claim", c)}*MIN(1,{a("enh_conv", c)}*{REF["m_mon"]})*{c}{UREF["surv"]}',
         NUM1)
urow_add('cat_spons', 'Category sponsorships (average)',
         lambda c, i: f'={REF["categories"]}*MIN(1,{a("cat_sell", c)}*{REF["m_mon"]})*{c}{UREF["surv"]}', NUM1)
urow_add('paying', 'Paying directory businesses (average)',
         lambda c, i: f'={c}{UREF["feat_spots"]}+{c}{UREF["enh_biz"]}+{c}{UREF["cat_spons"]}', NUM1,
         'Counts each paid product once; a business can hold more than one.')
usec('Revenue per launched town ($)')
urow_add('r_feat', 'Featured directory spots',
         lambda c, i: f'={c}{UREF["feat_spots"]}/3*{REF["feat_price"]}*12*(1-{REF["trial_loss"]})*{c}{UREF["live_frac"]}',
         USD, 'Spots x average price ($38 / 3) x 12, less free-trial months.')
urow_add('r_enh', 'Enhanced listings',
         lambda c, i: f'={c}{UREF["enh_biz"]}*{REF["enh_price"]}*12*{c}{UREF["live_frac"]}', USD)
urow_add('r_cat', 'Category sponsorships',
         lambda c, i: f'={c}{UREF["cat_spons"]}*{REF["cat_price"]}*12*{c}{UREF["live_frac"]}', USD)
urow_add('r_nl', 'Newsletter sponsorships',
         lambda c, i: f'={c}{UREF["subs"]}/1000*{REF["nl_cpm"]}*{REF["nl_sends"]}*MIN(1,{a("nl_sell", c)}*{REF["m_mon"]})*{c}{UREF["live_frac"]}',
         USD, 'Subscribers / 1,000 x CPM x issues x sell-through.')
urow_add('r_ban', 'Town sponsor banner',
         lambda c, i: f'={REF["banner"]}*12*MIN(1,{a("ban_sell", c)}*{REF["m_mon"]})*{c}{UREF["town_years"]}', USD)
urow_add('r_prog', 'Programmatic display',
         lambda c, i: f'={c}{UREF["pv"]}/1000*{REF["prog_rpm"]}*{REF["prog_fill"]}*{a("prog_on", c)}', USD)
urow_add('rev', 'Total revenue per launched town',
         lambda c, i: f'=SUM({c}{UREF["r_feat"]}:{c}{UREF["r_prog"]})', USD, bold=True, top=True)
urow_add('rev_dir', '  of which directory (self-serve, Stripe)',
         lambda c, i: f'={c}{UREF["r_feat"]}+{c}{UREF["r_enh"]}', USD)
urow_add('rev_direct', '  of which direct-sold (commissionable)',
         lambda c, i: f'={c}{UREF["r_cat"]}+{c}{UREF["r_nl"]}+{c}{UREF["r_ban"]}', USD)
usec('Direct costs per launched town ($)')
urow_add('c_ai', 'AI (writing and fact-checking)',
         lambda c, i: f'=MIN({REF["ai_cap"]},{REF["stories"]}*{REF["ai_story"]}+{REF["ai_overhead"]})*{REF["m_cost"]}*365*{c}{UREF["town_years"]}',
         USD, 'Stories a day x cost a story + overhead, capped at the daily cap.')
urow_add('c_host', 'Hosting, database and data',
         lambda c, i: f'={REF["hosting"]}*12*{REF["m_cost"]}*{c}{UREF["town_years"]}', USD)
urow_add('c_email', 'Email sending',
         lambda c, i: f'={c}{UREF["subs"]}*{REF["nl_sends"]}/1000*{REF["ses"]}*{c}{UREF["live_frac"]}', USD)
urow_add('c_stripe', 'Card processing',
         lambda c, i: f'={c}{UREF["rev_dir"]}*{REF["stripe_pct"]}+({c}{UREF["feat_spots"]}+{c}{UREF["enh_biz"]})*12*{REF["stripe_fixed"]}*{c}{UREF["live_frac"]}',
         USD)
urow_add('c_domain', 'Domain renewal', lambda c, i: '=0' if i == 0 else f'={REF["domain"]}*{c}{UREF["surv"]}', USD,
         'Bought at launch (launch costs); renewed from Age 2.')
urow_add('cogs', 'Total direct costs', lambda c, i: f'=SUM({c}{UREF["c_ai"]}:{c}{UREF["c_domain"]})', USD,
         bold=True, top=True)
urow_add('gp', 'Gross profit per launched town', lambda c, i: f'={c}{UREF["rev"]}-{c}{UREF["cogs"]}', USD,
         bold=True)
urow_add('gm', 'Gross margin', lambda c, i: f'=IF({c}{UREF["rev"]}=0,0,{c}{UREF["gp"]}/{c}{UREF["rev"]})', PCT)
usec('Town-level people and acquisition costs ($)')
urow_add('c_rev', 'Editorial standards (share)',
         lambda c, i: f'={REF["sal_rev"]}*(1+{REF["benefits"]})/{REF["per_reviewer"]}*{c}{UREF["town_years"]}', USD)
urow_add('c_ae', 'Advertiser partnerships (share) + commission',
         lambda c, i: f'={REF["sal_ae"]}*(1+{REF["benefits"]})/{REF["per_ae"]}*{c}{UREF["town_years"]}+{c}{UREF["rev_direct"]}*{REF["commission"]}',
         USD)
urow_add('c_sup', 'Support and success (share)',
         lambda c, i: f'={REF["sal_sup"]}*(1+{REF["benefits"]})/{REF["per_support"]}*{c}{UREF["town_years"]}', USD)
urow_add('c_cac', 'Paid subscriber acquisition',
         lambda c, i: f'={c}{UREF["new_subs"]}*{REF["sub_paid"]}*{REF["sub_cac"]}', USD)
urow_add('contrib', 'Town contribution (before central costs)',
         lambda c, i: f'={c}{UREF["gp"]}-SUM({c}{UREF["c_rev"]}:{c}{UREF["c_cac"]})', USD, bold=True, top=True)
urow_add('contrib_m', 'Contribution margin',
         lambda c, i: f'=IF({c}{UREF["rev"]}=0,0,{c}{UREF["contrib"]}/{c}{UREF["rev"]})', PCT)
urow_add('launch_cost', 'One-time launch cost (Age 1 only)',
         lambda c, i: (f'=({REF["launch_ai"]}+{REF["launch_domain"]}+{REF["launch_mkt"]})*{REF["m_cost"]}' if i == 0 else '=0'),
         USD)
urow_add('cum', 'Cumulative contribution after launch cost',
         lambda c, i: (f'={c}{UREF["contrib"]}-{c}{UREF["launch_cost"]}' if i == 0
                       else f'={COLS[i-1]}{urow}+{c}{UREF["contrib"]}-{c}{UREF["launch_cost"]}'),
         USD, 'Payback: the first age where this turns positive.', bold=True)
UREF['cum'] = urow - 1
U[f'A{urow+1}'] = ('Note: rows are expected values per launched town (survival-weighted), so a '
                   'surviving mature town earns slightly more than Age 5 shows.')
U[f'A{urow+1}'].font = F_NOTE

# ----------------------------------------------------------------------------------------------
# Monthly cash engine (bootstrapped): each month launches only what last month's cash and people
# allow, so the bank balance never drops below the floor. Every annual tab sums these months.
M = wb.create_sheet('Monthly')
title(M, 'Monthly cash engine (bootstrapped)',
      'Launches each month are limited by last month\'s cash and by staff the revenue can pay for. '
      'Founders are unpaid. Every annual tab sums these months.')
setw(M, {'A': 40, 'B': 8})
section(M, 4, 'Per launched town, per live month (from Unit Economics)')
header(M, 5, 'Town age', AGES)
MH = {}
mrow = 6
for key, label, kind in [
    ('surv', 'Share still live', 'stock'),
    ('r_feat', 'Featured directory spots', 'flow'), ('r_enh', 'Enhanced listings', 'flow'),
    ('r_cat', 'Category sponsorships', 'flow'), ('r_nl', 'Newsletter sponsorships', 'flow'),
    ('r_ban', 'Town sponsor banner', 'flow'), ('r_prog', 'Programmatic display', 'flow'),
    ('rev', 'Total revenue', 'flow'),
    ('c_ai', 'AI', 'flow'), ('c_host', 'Hosting, database and data', 'flow'), ('c_email', 'Email sending', 'flow'),
    ('c_stripe', 'Card processing', 'flow'), ('c_domain', 'Domain renewal', 'flow'),
    ('c_cac', 'Paid subscriber acquisition', 'flow'),
    ('subs', 'Newsletter subscribers', 'stock'), ('pv', 'Pageviews', 'flow'),
    ('paying', 'Paid directory products', 'stock'),
]:
    M[f'A{mrow}'] = label
    for c in COLS:
        src = f"'Unit Economics'!{c}{UREF[key]}"
        M[f'{c}{mrow}'] = f'={src}' if kind == 'stock' else f"={src}/(12*'Unit Economics'!{c}{UREF['live_frac']})"
        M[f'{c}{mrow}'].number_format = PCT if key == 'surv' else USD2
    MH[key] = mrow
    mrow += 1
M[f'A{mrow}'] = 'Age 1 ramp: share of the Age 1 monthly value, by month of life (1-12)'
for k in range(12):
    cell = M.cell(row=mrow, column=3 + k, value=round(k / 5.5, 4))
    cell.font = F_INPUT
    cell.number_format = '0.00'
M[f'O{mrow}'] = ('Estimate: revenue and audience start at zero (free trials, no readers yet) and grow '
                 'evenly through the launch year, averaging the Age 1 value. Costs start in full.')
M[f'O{mrow}'].font = F_NOTE
MH['ramp'] = mrow
mrow += 1
M[f'A{mrow}'] = 'One-time launch cost'
M[f'C{mrow}'] = f"='Unit Economics'!C{UREF['launch_cost']}"
MH['launch_cost'] = mrow
mrow += 1
M[f'A{mrow}'] = 'New town cash burn, by month of life (1-12)'
for k in range(12):
    rc = get_column_letter(3 + k)
    M[f'{rc}{mrow}'] = (f'=MAX(0,SUM($C${MH["c_ai"]}:$C${MH["c_host"]})+$C${MH["c_domain"]}+{REF["legal_town"]}/12'
                        f'+{rc}{MH["ramp"]}*(SUM($C${MH["c_email"]}:$C${MH["c_stripe"]})+$C${MH["c_cac"]}-$C${MH["rev"]})'
                        + ('' if k == 0 else
                           f'+({rc}{MH["ramp"]}-{get_column_letter(2 + k)}{MH["ramp"]})*$C${MH["rev"]}*12*{REF["nwc"]}')
                        + ')')
    M[f'{rc}{mrow}'].number_format = USD2
MH['burn'] = mrow
mrow += 1
M[f'A{mrow}'] = 'Burn still ahead after month of life (1-12)'
for k in range(12):
    rc = get_column_letter(3 + k)
    M[f'{rc}{mrow}'] = '=0' if k == 11 else f'=SUM({get_column_letter(4 + k)}{MH["burn"]}:N{MH["burn"]})'
    M[f'{rc}{mrow}'].number_format = USD2
MH['ahead'] = mrow
mrow += 1
M[f'A{mrow}'] = 'Cash set aside per launch'
M[f'C{mrow}'] = f'=C{MH["launch_cost"]}+SUM(C{MH["burn"]}:N{MH["burn"]})'
M[f'D{mrow}'] = 'Launch cost plus every month of burn until the new town pays its own way.'
M[f'D{mrow}'].font = F_NOTE
MH['per_launch'] = mrow
M[f'A{mrow + 1}'] = 'Fixed costs a month, FY2027 (platform and fixed insurance)'
M[f'C{mrow + 1}'] = f'={REF["platform_base"]}+Assumptions!C${REF["legal_fixed"]}/12'
MH['fixed'] = mrow + 1
for r in (MH['launch_cost'], MH['per_launch'], MH['fixed']):
    M[f'C{r}'].number_format = USD2
    M[f'C{r}'].font = F_BOLD
mrow += 3

MKEYS = [
    ('m', 'Month'), ('fy', 'Fiscal year'), ('yi', 'Year index'),
    ('cap', 'Launch capacity'), ('pool', 'Towns left to launch'), ('reserve', 'Cash held for young towns'),
    ('avail', 'Cash free for launches'),
    ('launch', 'Towns launched'), ('cum', 'Cumulative launched'),
    ('a1', 'Launched, Age 1'), ('r1', 'Age 1, ramp-weighted'), ('a2', 'Launched, Age 2'), ('a3', 'Launched, Age 3'), ('a4', 'Launched, Age 4'),
    ('a5', 'Launched, Age 5'), ('live', 'Live towns'),
    ('r_feat', 'Featured spots'), ('r_enh', 'Enhanced'), ('r_cat', 'Category sponsors'), ('r_nl', 'Newsletter'),
    ('r_ban', 'Banners'), ('r_prog', 'Programmatic'), ('rev', 'Revenue'),
    ('c_ai', 'AI'), ('c_host', 'Hosting'), ('c_email', 'Email'), ('c_stripe', 'Card fees'),
    ('c_domain', 'Domain renewals'), ('c_newdom', 'New domains'), ('cogs', 'Cost of revenue'),
    ('subs', 'Subscribers'), ('pv', 'Pageviews'), ('paying', 'Paid products'),
    ('n_rev', 'FTE needed: editorial'), ('n_ae', 'FTE needed: partnerships'), ('n_sup', 'FTE needed: support'),
    ('n_ops', 'FTE needed: launches'), ('n_lead', 'FTE wanted: leadership'), ('n_eng', 'FTE wanted: engineering'),
    ('n_ga', 'FTE wanted: G&A'), ('n_fte', 'FTE needed, total'),
    ('k_rev', 'Cost if paid: editorial'), ('k_ae', 'Cost if paid: partnerships'), ('k_sup', 'Cost if paid: support'),
    ('k_ops', 'Cost if paid: launches'), ('k_lead', 'Cost if paid: leadership'), ('k_eng', 'Cost if paid: engineering'),
    ('k_ga', 'Cost if paid: G&A'), ('k_cost', 'Cost if all paid'),
    ('budget', 'Staff budget'), ('paid_ops', 'Paid: town work'), ('f', 'Share of town work paid'),
    ('paid_c', 'Paid: central hires'), ('fc', 'Share of central hires paid'), ('paid', 'Paid staff cost'),
    ('paid_fte', 'Paid FTE'), ('unfunded', 'Unfunded town work (FTE)'),
    ('p_rev', 'Paid: editorial'), ('p_ae', 'Paid: partnerships'), ('p_sup', 'Paid: support'), ('p_ops', 'Paid: launches'),
    ('p_lead', 'Paid: leadership'), ('p_eng', 'Paid: engineering'), ('p_ga', 'Paid: G&A'),
    ('launch_mkt', 'Launch marketing and AI'), ('sub_cac', 'Paid subscribers'), ('brand', 'Corporate marketing'),
    ('legal', 'Legal and insurance'), ('tools', 'Platform, software and tools'), ('opex', 'Operating expenses'),
    ('ebitda', 'EBITDA'), ('capex', 'Equipment'), ('tax', 'Income tax paid'), ('nwc', 'Working capital'),
    ('equity', 'Equity raised'), ('flow', 'Net cash flow'), ('close', 'Closing cash'),
]
MC = {k: get_column_letter(i + 1) for i, (k, _) in enumerate(MKEYS)}
T0 = mrow  # header row of the monthly table
M0 = T0 + 1  # month 1
MONTHS = 60
for k, label in MKEYS:
    cell = M[f'{MC[k]}{T0}']
    cell.value = label
    cell.font = F_HEAD
    cell.fill = FILL_HEAD
    cell.alignment = Alignment(wrap_text=True, horizontal='right')
    if k not in ('m', 'fy', 'yi'):
        M.column_dimensions[MC[k]].width = 13
M.column_dimensions['A'].width = 40
ROLE_KEYS = ['rev', 'ae', 'sup', 'ops', 'lead', 'eng', 'ga']
ROLE_SAL = {'rev': 'sal_rev', 'ae': 'sal_ae', 'sup': 'sal_sup', 'ops': 'sal_rev', 'lead': 'sal_lead',
            'eng': 'sal_eng', 'ga': 'sal_ga'}


def mref(k, r):
    return f'{MC[k]}{r}'


RAMPED = ('r_feat', 'r_enh', 'r_cat', 'r_nl', 'r_ban', 'r_prog', 'c_email', 'c_stripe', 'c_cac', 'subs', 'pv',
          'paying')


def by_age(key, r):
    """Sum over ages of launched-town counts x the per-month value at that age. Revenue and audience
    use the ramp-weighted Age 1 count; costs start in full."""
    first = 'r1' if key in RAMPED else 'a1'
    terms = [f'{mref(first, r)}*$C${MH[key]}']
    terms += [f'{mref("a" + str(n + 1), r)}*${COLS[n]}${MH[key]}' for n in range(1, 5)]
    return '+'.join(terms)


def fy_rows(j):
    return M0 + 12 * j, M0 + 12 * j + 11


for i in range(MONTHS):
    r = M0 + i
    m = i + 1
    yi = i // 12
    p = r - 1
    f = {}
    f['m'] = m
    f['fy'] = int(YEARS[yi])
    f['yi'] = yi
    f['cap'] = f'=ROUND(Assumptions!{COLS[yi]}${REF["launch_cap"]}*{REF["m_launch"]}/12,0)'
    f['pool'] = '=Rollout!$C$10' + ('' if m == 1 else f'-{mref("cum", p)}')
    prev_close = REF['open_cash'] if m == 1 else mref('close', p)
    f['reserve'] = '=' + '+'.join(f'{mref("launch", r - k)}*${get_column_letter(3 + k)}${MH["ahead"]}'
                                  for k in range(min(12, m)))
    held = '' if m == 1 else f'-{mref("reserve", p)}'
    fixed = (f'{REF["platform_base"]}+Assumptions!{COLS[yi]}${REF["legal_fixed"]}/12'
             f'+{("0" if m == 1 else mref("paid", p))}')
    f['avail'] = f'={prev_close}-{REF["cash_floor"]}{held}-{REF["buffer_months"]}*({fixed})'
    fit = f'MAX(0,MIN({mref("cap", r)},{mref("pool", r)},INT({mref("avail", r)}/$C${MH["per_launch"]})))'
    f['launch'] = f'={fit}' if m == 1 else f'=IF({mref("unfunded", p)}>0,0,{fit})'
    f['cum'] = f'={mref("launch", r)}' if m == 1 else f'={mref("cum", p)}+{mref("launch", r)}'
    for n in range(5):
        lo, hi = max(1, m - 12 * (n + 1) + 1), m - 12 * n
        f[f'a{n + 1}'] = f'=SUM({mref("launch", M0 + lo - 1)}:{mref("launch", M0 + hi - 1)})' if hi >= 1 else '=0'
    f['r1'] = '=' + '+'.join(f'{mref("launch", r - k)}*${get_column_letter(3 + k)}${MH["ramp"]}'
                             for k in range(min(12, m)))
    f['live'] = '=' + by_age('surv', r)
    for k in ('r_feat', 'r_enh', 'r_cat', 'r_nl', 'r_ban', 'r_prog', 'c_ai', 'c_host', 'c_email', 'c_stripe',
              'c_domain', 'subs', 'pv', 'paying'):
        f[k] = '=' + by_age(k, r)
    f['rev'] = f'=SUM({mref("r_feat", r)}:{mref("r_prog", r)})'
    f['c_newdom'] = f'={mref("launch", r)}*{REF["launch_domain"]}*{REF["m_cost"]}'
    f['cogs'] = f'=SUM({mref("c_ai", r)}:{mref("c_newdom", r)})'
    f['n_rev'] = f'={mref("live", r)}/{REF["per_reviewer"]}'
    f['n_ae'] = f'={mref("live", r)}/{REF["per_ae"]}'
    f['n_sup'] = f'={mref("live", r)}/{REF["per_support"]}'
    f['n_ops'] = f'={mref("launch", r)}/({REF["per_launcher"]}/12)'
    for k, s in (('lead', 'hc_lead'), ('eng', 'hc_eng'), ('ga', 'hc_ga')):
        f['n_' + k] = f'=Assumptions!{COLS[yi]}${REF[s]}'
    f['n_fte'] = f'=SUM({mref("n_rev", r)}:{mref("n_ga", r)})'
    for k in ROLE_KEYS:
        f['k_' + k] = f'={mref("n_" + k, r)}*{REF[ROLE_SAL[k]]}*(1+{REF["benefits"]})*(1+{REF["raise"]})^{yi}/12'
    f['k_cost'] = f'=SUM({mref("k_rev", r)}:{mref("k_ga", r)})'
    f['budget'] = ('=0' if m == 1 else
                   f'=MIN({REF["people_share"]}*{mref("rev", p)},MAX(0,{mref("ebitda", p)}+{mref("paid", p)}))')
    # Town work (editorial, partnerships, support, launches) is paid first, beyond founder time;
    # central hires (leadership, engineering, G&A) only from what the budget has left.
    n_town = f'SUM({mref("n_rev", r)}:{mref("n_ops", r)})'
    k_town = f'SUM({mref("k_rev", r)}:{mref("k_ops", r)})'
    k_central = f'SUM({mref("k_lead", r)}:{mref("k_ga", r)})'
    f['paid_ops'] = (f'=IF({n_town}<=0,0,MIN({mref("budget", r)},'
                     f'{k_town}*MAX(0,1-{REF["founder_fte"]}/{n_town})))')
    f['f'] = f'=IF({k_town}=0,0,{mref("paid_ops", r)}/{k_town})'
    f['paid_c'] = f'=MIN({mref("budget", r)}-{mref("paid_ops", r)},{k_central})'
    f['fc'] = f'=IF({k_central}=0,0,{mref("paid_c", r)}/{k_central})'
    f['paid'] = f'={mref("paid_ops", r)}+{mref("paid_c", r)}'
    f['paid_fte'] = (f'={n_town}*{mref("f", r)}'
                     f'+SUM({mref("n_lead", r)}:{mref("n_ga", r)})*{mref("fc", r)}')
    f['unfunded'] = f'=MAX(0,{n_town}*(1-{mref("f", r)})-{REF["founder_fte"]})'
    for k in ROLE_KEYS:
        share = 'f' if k in ('rev', 'ae', 'sup', 'ops') else 'fc'
        f['p_' + k] = f'={mref("k_" + k, r)}*{mref(share, r)}'
    f['launch_mkt'] = f'={mref("launch", r)}*({REF["launch_mkt"]}+{REF["launch_ai"]})*{REF["m_cost"]}'
    f['sub_cac'] = '=' + by_age('c_cac', r)
    f['brand'] = f'={REF["brand_share"]}*{mref("rev", r)}'
    f['legal'] = f'=Assumptions!{COLS[yi]}${REF["legal_fixed"]}/12+{mref("live", r)}*{REF["legal_town"]}/12'
    f['tools'] = f'={mref("paid_fte", r)}*{REF["tools"]}/12+{REF["platform_base"]}'
    f['opex'] = f'=SUM({mref("p_rev", r)}:{mref("tools", r)})'
    f['ebitda'] = f'={mref("rev", r)}-{mref("cogs", r)}-{mref("opex", r)}'
    f['capex'] = (f'={mref("paid_fte", r)}*{REF["laptop"]}' if m == 1
                  else f'=MAX(0,{mref("paid_fte", r)}-{mref("paid_fte", p)})*{REF["laptop"]}')
    f['tax'] = f"=-'P&L'!{COLS[yi - 1]}@TAX@" if m % 12 == 1 and m > 1 else '=0'
    f['nwc'] = (f'=-{mref("rev", r)}*12*{REF["nwc"]}' if m == 1
                else f'=-({mref("rev", r)}-{mref("rev", p)})*12*{REF["nwc"]}')
    f['equity'] = f'=Assumptions!{COLS[yi]}${REF["equity"]}' if m % 12 == 1 else '=0'
    f['flow'] = f'={mref("ebitda", r)}-{mref("capex", r)}+{mref("tax", r)}+{mref("nwc", r)}+{mref("equity", r)}'
    f['close'] = f'={prev_close}+{mref("flow", r)}'
    for k, _ in MKEYS:
        cell = M[f'{MC[k]}{r}']
        cell.value = f[k]
        if k in ('m', 'fy', 'yi'):
            cell.number_format = '0'
        elif k in ('f', 'fc'):
            cell.number_format = PCT
        elif k.startswith('n_') or k in ('paid_fte', 'unfunded'):
            cell.number_format = '0.00'
        elif k in ('cap', 'pool', 'launch', 'cum', 'live', 'subs', 'pv', 'paying', 'r1') or k[0] == 'a' and k[1:].isdigit():
            cell.number_format = NUM
        else:
            cell.number_format = USD
    M[f'{MC["close"]}{r}'].font = F_BOLD
M.freeze_panes = f'D{M0}'


def msum(k, j):
    a_, b_ = fy_rows(j)
    return f'SUM(Monthly!{MC[k]}{a_}:{MC[k]}{b_})'


def mend(k, j):
    return f'Monthly!{MC[k]}{fy_rows(j)[1]}'


# Rollout (annual view of the monthly launches)
title(R, 'Town rollout', 'Launches by year and live towns, from the Monthly sheet.')
setw(R, {'A': 46, 'B': 4, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 60})
header(R, 4, 'Fiscal year', YEARS)
R['A5'] = 'New towns launched'
R['A6'] = 'Cumulative towns launched'
R['A7'] = 'Live towns at year end'
R['A8'] = 'Average live towns in year'
R['A9'] = 'Launch capacity not used (cash or people)'
for j, c in enumerate(COLS):
    R[f'{c}5'] = f'={msum("launch", j)}'
    R[f'{c}6'] = f'={mend("cum", j)}'
    R[f'{c}7'] = f'={mend("live", j)}'
    a_, b_ = fy_rows(j)
    R[f'{c}8'] = f'=AVERAGE(Monthly!{MC["live"]}{a_}:{MC["live"]}{b_})'
    R[f'{c}9'] = f'=MAX(0,MIN({msum("cap", j)},$C$10-{c}6+{c}5)-{c}5)'
    for r in (5, 6, 7, 8, 9):
        R[f'{c}{r}'].number_format = NUM
R['H5'] = 'Each month launches what last month\'s cash and people allow (Monthly sheet).'
R['H9'] = 'Capacity the team has but cash or unpaid work holds back.'
R['A10'] = 'Launchable towns (pool)'
R['C10'] = (f'=ROUND(MIN({REF["addressable"]},{REF["addressable"]}*{REF["ready_share"]}'
            f'*{REF["name_share"]}*{REF["m_launch"]}),0)')
R['C10'].number_format = NUM
R['C10'].font = F_BOLD
R['H10'] = 'Towns over 2,000 people x readiness pass rate x naming rule (x launch multiplier).'
R['A11'] = 'Pool left at year end'
for c in COLS:
    R[f'{c}11'] = f'=$C$10-{c}6'
    R[f'{c}11'].number_format = NUM
for r in range(5, 12):
    R[f'A{r}'].font = F_BASE

# Revenue
V = wb.create_sheet('Revenue')
title(V, 'Revenue build ($)', 'Network totals: the sum of each fiscal year\'s months on the Monthly sheet.')
setw(V, {'A': 46, 'B': 4, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 40})
header(V, 4, 'Fiscal year', YEARS)
vrow = 5
VREF = {}
for key, label in [('r_feat', 'Featured directory spots'), ('r_enh', 'Enhanced listings'),
                   ('r_cat', 'Category sponsorships'), ('r_nl', 'Newsletter sponsorships'),
                   ('r_ban', 'Town sponsor banners'), ('r_prog', 'Programmatic display')]:
    V[f'A{vrow}'] = label
    for j, c in enumerate(COLS):
        V[f'{c}{vrow}'] = f'={msum(key, j)}'
        V[f'{c}{vrow}'].number_format = USD
    VREF[key] = vrow
    vrow += 1
V[f'A{vrow}'] = 'Total revenue'
V[f'A{vrow}'].font = F_BOLD
for c in COLS:
    V[f'{c}{vrow}'] = f'=SUM({c}5:{c}{vrow-1})'
    V[f'{c}{vrow}'].number_format = USD
    V[f'{c}{vrow}'].font = F_BOLD
    V[f'{c}{vrow}'].border = TOP
VREF['total'] = vrow
vrow += 1
V[f'A{vrow}'] = 'Growth'
for i, c in enumerate(COLS):
    V[f'{c}{vrow}'] = '=0' if i == 0 else f'=IF({COLS[i-1]}{VREF["total"]}=0,0,{c}{VREF["total"]}/{COLS[i-1]}{VREF["total"]}-1)'
    V[f'{c}{vrow}'].number_format = PCT
vrow += 2
V[f'A{vrow}'] = 'Revenue mix'
V[f'A{vrow}'].font = F_BOLD
vrow += 1
for key, label in [('r_feat', 'Featured spots'), ('r_enh', 'Enhanced listings'), ('r_cat', 'Category sponsorships'),
                   ('r_nl', 'Newsletter'), ('r_ban', 'Banners'), ('r_prog', 'Programmatic')]:
    V[f'A{vrow}'] = label
    for c in COLS:
        V[f'{c}{vrow}'] = f'=IF({c}${VREF["total"]}=0,0,{c}{VREF[key]}/{c}${VREF["total"]})'
        V[f'{c}{vrow}'].number_format = PCT
    vrow += 1

# Headcount
H = wb.create_sheet('Headcount')
title(H, 'Headcount and people costs', 'Founders are unpaid. Work beyond founder time is paid only out of '
      'revenue (Assumptions: staff budget); the rest is unfunded and holds launches.')
setw(H, {'A': 46, 'B': 4, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 50})
header(H, 4, 'Fiscal year', YEARS)
roles = [
    ('lead', 'Leadership (not founders)'),
    ('eng', 'Engineering and product'),
    ('ga', 'Finance, legal and operations'),
    ('rev', 'Editorial standards'),
    ('ae', 'Advertiser partnerships'),
    ('sup', 'Support and success'),
    ('ops', 'Launch operations'),
]
H['A5'] = 'FTE needed (year end)'
H['A5'].font = F_BOLD
hrow = 6
HREF = {}
for key, label in roles:
    H[f'A{hrow}'] = label
    for j, c in enumerate(COLS):
        H[f'{c}{hrow}'] = f'={mend("n_" + key, j)}'
        H[f'{c}{hrow}'].number_format = NUM1
    HREF['need_' + key] = hrow
    hrow += 1
for key, label, src in [('need', 'Total FTE needed', 'n_fte'), ('paid_fte', 'Paid staff (FTE)', 'paid_fte'),
                        ('unfunded', 'Unfunded town work, FTE (holds launches)', 'unfunded')]:
    H[f'A{hrow}'] = label
    H[f'A{hrow}'].font = F_BOLD
    for j, c in enumerate(COLS):
        H[f'{c}{hrow}'] = f'={mend(src, j)}'
        H[f'{c}{hrow}'].number_format = NUM1
        H[f'{c}{hrow}'].font = F_BOLD
    HREF[key] = hrow
    hrow += 1
H[f'A{hrow}'] = 'Founders\' time (unpaid FTE)'
for c in COLS:
    H[f'{c}{hrow}'] = f'={REF["founder_fte"]}'
    H[f'{c}{hrow}'].number_format = NUM1
    H[f'{c}{hrow}'].font = F_LINK
HREF['founders'] = hrow
HREF['total'] = HREF['paid_fte']
hrow += 2
H[f'A{hrow}'] = 'Paid people cost ($)'
H[f'A{hrow}'].font = F_BOLD
hrow += 1
for key, label in roles:
    H[f'A{hrow}'] = label
    for j, c in enumerate(COLS):
        H[f'{c}{hrow}'] = f'={msum("p_" + key, j)}'
        H[f'{c}{hrow}'].number_format = USD
    HREF['cost_' + key] = hrow
    hrow += 1
H[f'A{hrow}'] = 'Total people cost'
H[f'A{hrow}'].font = F_BOLD
first_cost = HREF['cost_lead']
for c in COLS:
    H[f'{c}{hrow}'] = f'=SUM({c}{first_cost}:{c}{hrow-1})'
    H[f'{c}{hrow}'].number_format = USD
    H[f'{c}{hrow}'].font = F_BOLD
    H[f'{c}{hrow}'].border = TOP
HREF['cost_total'] = hrow
H[f'A{hrow+2}'] = ('Each month, needed work beyond founder time is paid up to the staff budget (a share of '
                   'last month\'s revenue), split across roles in proportion to their cost.')
H[f'A{hrow+2}'].font = F_NOTE

# Costs
C = wb.create_sheet('Costs')
title(C, 'Costs ($)', 'Sums of the Monthly sheet; people from Headcount.')
setw(C, {'A': 46, 'B': 4, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 40})
header(C, 4, 'Fiscal year', YEARS)
crow = 5
CREF = {}


def crow_add(name, label, fn, bold=False, top=False, fmt=USD, font=None):
    global crow
    C[f'A{crow}'] = label
    C[f'A{crow}'].font = F_BOLD if bold else F_BASE
    for j, c in enumerate(COLS):
        C[f'{c}{crow}'] = fn(c, j)
        C[f'{c}{crow}'].number_format = fmt
        if bold:
            C[f'{c}{crow}'].font = F_BOLD
        if font:
            C[f'{c}{crow}'].font = font
        if top:
            C[f'{c}{crow}'].border = TOP
    CREF[name] = crow
    crow += 1


section(C, crow, 'Cost of revenue')
crow += 1
crow_add('ai', 'AI (writing and fact-checking)', lambda c, j: f'={msum("c_ai", j)}')
crow_add('host', 'Hosting, database and data', lambda c, j: f'={msum("c_host", j)}')
crow_add('email', 'Email sending', lambda c, j: f'={msum("c_email", j)}')
crow_add('stripe', 'Card processing', lambda c, j: f'={msum("c_stripe", j)}')
crow_add('domain', 'Domains (renewals and new)', lambda c, j: f'={msum("c_domain", j)}+{msum("c_newdom", j)}')
crow_add('cogs', 'Total cost of revenue', lambda c, j: f'=SUM({c}{CREF["ai"]}:{c}{CREF["domain"]})', bold=True,
         top=True)
crow += 1
section(C, crow, 'Operating expenses')
crow += 1
crow_add('p_editorial', 'Editorial standards (people)', lambda c, j: f'=Headcount!{c}{HREF["cost_rev"]}', font=F_LINK)
crow_add('p_sales', 'Advertiser partnerships (people)', lambda c, j: f'=Headcount!{c}{HREF["cost_ae"]}', font=F_LINK)
crow_add('commission', 'Sales commission',
         lambda c, j: f'=(Revenue!{c}{VREF["r_cat"]}+Revenue!{c}{VREF["r_nl"]}+Revenue!{c}{VREF["r_ban"]})*{REF["commission"]}')
crow_add('p_support', 'Support and success (people)', lambda c, j: f'=Headcount!{c}{HREF["cost_sup"]}', font=F_LINK)
crow_add('p_ops', 'Launch operations (people)', lambda c, j: f'=Headcount!{c}{HREF["cost_ops"]}', font=F_LINK)
crow_add('launch_mkt', 'Launch marketing and launch AI', lambda c, j: f'={msum("launch_mkt", j)}')
crow_add('sub_cac', 'Paid subscriber acquisition', lambda c, j: f'={msum("sub_cac", j)}')
crow_add('brand', 'Corporate marketing', lambda c, j: f'={msum("brand", j)}')
crow_add('p_eng', 'Engineering and product (people)', lambda c, j: f'=Headcount!{c}{HREF["cost_eng"]}', font=F_LINK)
crow_add('p_lead', 'Leadership (people, not founders)', lambda c, j: f'=Headcount!{c}{HREF["cost_lead"]}', font=F_LINK)
crow_add('p_ga', 'Finance, legal and operations (people)', lambda c, j: f'=Headcount!{c}{HREF["cost_ga"]}',
         font=F_LINK)
crow_add('legal', 'Legal and media liability insurance', lambda c, j: f'={msum("legal", j)}')
crow_add('tools', 'Platform, software and tools', lambda c, j: f'={msum("tools", j)}')
crow_add('opex', 'Total operating expenses', lambda c, j: f'=SUM({c}{CREF["p_editorial"]}:{c}{CREF["tools"]})',
         bold=True, top=True)

# P&L
P = wb.create_sheet('P&L')
title(P, 'Income statement ($)', 'Fiscal years 2027-2031.')
setw(P, {'A': 46, 'B': 4, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 40})
header(P, 4, 'Fiscal year', YEARS)
prow = 5
PREF = {}


def prow_add(name, label, fn, fmt=USD, bold=False, top=False, font=None):
    global prow
    P[f'A{prow}'] = label
    P[f'A{prow}'].font = F_BOLD if bold else F_BASE
    for j, c in enumerate(COLS):
        P[f'{c}{prow}'] = fn(c, j)
        P[f'{c}{prow}'].number_format = fmt
        if font:
            P[f'{c}{prow}'].font = font
        if bold:
            P[f'{c}{prow}'].font = F_BOLD
        if top:
            P[f'{c}{prow}'].border = TOP
    PREF[name] = prow
    prow += 1


prow_add('rev', 'Revenue', lambda c, j: f'=Revenue!{c}{VREF["total"]}', bold=True)
prow_add('cogs', 'Cost of revenue', lambda c, j: f'=Costs!{c}{CREF["cogs"]}', font=F_LINK)
prow_add('gp', 'Gross profit', lambda c, j: f'={c}{PREF["rev"]}-{c}{PREF["cogs"]}', bold=True, top=True)
prow_add('gm', 'Gross margin', lambda c, j: f'=IF({c}{PREF["rev"]}=0,0,{c}{PREF["gp"]}/{c}{PREF["rev"]})', fmt=PCT)
prow += 1
prow_add('editorial', 'Editorial', lambda c, j: f'=Costs!{c}{CREF["p_editorial"]}', font=F_LINK)
prow_add('sm', 'Sales and marketing',
         lambda c, j: f'=Costs!{c}{CREF["p_sales"]}+Costs!{c}{CREF["commission"]}+Costs!{c}{CREF["p_support"]}+Costs!{c}{CREF["p_ops"]}'
                      f'+Costs!{c}{CREF["launch_mkt"]}+Costs!{c}{CREF["sub_cac"]}+Costs!{c}{CREF["brand"]}',
         font=F_LINK)
prow_add('rd', 'Engineering and product', lambda c, j: f'=Costs!{c}{CREF["p_eng"]}', font=F_LINK)
prow_add('ga', 'General and administrative',
         lambda c, j: f'=Costs!{c}{CREF["p_lead"]}+Costs!{c}{CREF["p_ga"]}+Costs!{c}{CREF["legal"]}+Costs!{c}{CREF["tools"]}',
         font=F_LINK)
prow_add('opex', 'Total operating expenses', lambda c, j: f'=SUM({c}{PREF["editorial"]}:{c}{PREF["ga"]})', top=True)
prow += 1
prow_add('ebitda', 'EBITDA', lambda c, j: f'={c}{PREF["gp"]}-{c}{PREF["opex"]}', bold=True)
prow_add('ebitda_m', 'EBITDA margin',
         lambda c, j: f'=IF({c}{PREF["rev"]}=0,0,{c}{PREF["ebitda"]}/{c}{PREF["rev"]})', fmt=PCT)
prow_add('capex', 'Capital expenditure (memo)', lambda c, j: f'={msum("capex", j)}')
prow_add('da', 'Depreciation',
         lambda c, j: '=(' + '+'.join(f'{COLS[k]}{prow-1}' for k in range(j + 1)) + f')/{REF["dep_years"]}'
         if j < 3 else '=(' + '+'.join(f'{COLS[k]}{prow-1}' for k in range(j - 2, j + 1)) + f')/{REF["dep_years"]}')
prow_add('ebit', 'Operating income (EBIT)', lambda c, j: f'={c}{PREF["ebitda"]}-{c}{PREF["da"]}', bold=True)
prow_add('nol_open', 'Tax losses available, opening',
         lambda c, j: '=0' if j == 0 else f'={COLS[j-1]}{prow+2}')
prow_add('taxable', 'Taxable income after losses used',
         lambda c, j: f'=MAX(0,{c}{PREF["ebit"]}-{c}{PREF["nol_open"]})')
prow_add('nol_close', 'Tax losses available, closing',
         lambda c, j: f'=MAX(0,{c}{PREF["nol_open"]}-MAX(0,{c}{PREF["ebit"]}))+MAX(0,-{c}{PREF["ebit"]})')
prow_add('tax', 'Income tax', lambda c, j: f'={c}{PREF["taxable"]}*{REF["tax"]}')
prow_add('ni', 'Net income', lambda c, j: f'={c}{PREF["ebit"]}-{c}{PREF["tax"]}', bold=True, top=True)
P[f'A{prow+1}'] = 'Income tax is paid in the first month of the next fiscal year (Monthly sheet).'
P[f'A{prow+1}'].font = F_NOTE
for i in range(MONTHS):
    cell = M[f'{MC["tax"]}{M0 + i}']
    cell.value = cell.value.replace('@TAX@', str(PREF['tax']))

# Cash flow
CF = wb.create_sheet('Cash Flow')
title(CF, 'Cash flow and funding ($)', 'Bootstrapped: sums of the Monthly sheet, no outside money unless '
      'Assumptions adds equity.')
setw(CF, {'A': 46, 'B': 4, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 50})
header(CF, 4, 'Fiscal year', YEARS)
CF['A5'] = 'EBITDA'
CF['A6'] = 'Income tax paid'
CF['A7'] = 'Capital expenditure'
CF['A8'] = 'Increase in net working capital'
CF['A9'] = 'Free cash flow'
CF['A10'] = 'Equity raised'
CF['A11'] = 'Net change in cash'
CF['A12'] = 'Opening cash'
CF['A13'] = 'Closing cash'
CF['A14'] = 'Cash below zero in any month? (1 = yes: fix the plan)'
for j, c in enumerate(COLS):
    prev = COLS[j - 1] if j else None
    a_, b_ = fy_rows(j)
    CF[f'{c}5'] = f"='P&L'!{c}{PREF['ebitda']}"
    CF[f'{c}6'] = f'={msum("tax", j)}'
    CF[f'{c}7'] = f"=-'P&L'!{c}{PREF['capex']}"
    CF[f'{c}8'] = f'={msum("nwc", j)}'
    CF[f'{c}9'] = f'=SUM({c}5:{c}8)'
    CF[f'{c}10'] = f'={msum("equity", j)}'
    CF[f'{c}11'] = f'={c}9+{c}10'
    CF[f'{c}12'] = f'={REF["open_cash"]}' if j == 0 else f'={prev}13'
    CF[f'{c}13'] = f'={c}12+{c}11'
    CF[f'{c}14'] = f'=IF(MIN(Monthly!{MC["close"]}{a_}:{MC["close"]}{b_})<0,1,0)'
    for r in range(5, 14):
        CF[f'{c}{r}'].number_format = USD
    CF[f'{c}14'].number_format = '0'
    for r in (5, 7):
        CF[f'{c}{r}'].font = F_LINK
    for r in (9, 13):
        CF[f'{c}{r}'].font = F_BOLD
        CF[f'{c}{r}'].border = TOP
for r in (9, 13):
    CF[f'A{r}'].font = F_BOLD
CF['A16'] = 'Lowest month-end cash, FY2027-FY2031'
CF['C16'] = f'=MIN(Monthly!{MC["close"]}{M0}:{MC["close"]}{M0 + MONTHS - 1})'
CF['C16'].number_format = USD
CF['C16'].font = F_BOLD
CF['H16'] = 'Must stay at or above the cash floor in Assumptions.'
CF['H16'].font = F_NOTE
CF['A17'] = 'Cumulative free cash flow before turning positive (funding need)'
CF['C17'] = '=-MIN(0,C9,C9+D9,C9+D9+E9,C9+D9+E9+F9,C9+D9+E9+F9+G9)'
CF['C17'].number_format = USD
CF['C17'].font = F_BOLD
CF['H17'] = 'Deepest point of cumulative yearly free cash flow; bootstrapped, opening cash covers it.'
CF['H17'].font = F_NOTE

# KPIs
K = wb.create_sheet('KPIs')
title(K, 'Key metrics', 'Operating and financial KPIs, all formulas.')
setw(K, {'A': 50, 'B': 4, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 50})
header(K, 4, 'Fiscal year', YEARS)
krow = 5
KREF = {}


def krow_add(name, label, fn, fmt, note=''):
    global krow
    K[f'A{krow}'] = label
    for j, c in enumerate(COLS):
        K[f'{c}{krow}'] = fn(c, j)
        K[f'{c}{krow}'].number_format = fmt
    if note:
        K[f'H{krow}'] = note
        K[f'H{krow}'].font = F_NOTE
    KREF[name] = krow
    krow += 1


section(K, krow, 'Footprint and audience')
krow += 1
krow_add('launch', 'Towns launched in year', lambda c, j: f'=Rollout!{c}5', NUM)
krow_add('live', 'Live towns at year end', lambda c, j: f'=Rollout!{c}7', NUM)
krow_add('subs', 'Newsletter subscribers (year end)', lambda c, j: f'={mend("subs", j)}', NUM)
krow_add('pv', 'Pageviews in year', lambda c, j: f'={msum("pv", j)}', NUM)
section(K, krow, 'Customers')
krow += 1
krow_add('paying', 'Paid directory products (average in year)',
         lambda c, j: f'={msum("paying", j)}/12', NUM)
krow_add('arpb', 'Directory revenue per paid product per month ($)',
         lambda c, j: f'=IF({c}{KREF["paying"]}=0,0,(Revenue!{c}{VREF["r_feat"]}+Revenue!{c}{VREF["r_enh"]}+Revenue!{c}{VREF["r_cat"]})/{c}{KREF["paying"]}/12)',
         USD2)
section(K, krow, 'Financial')
krow += 1
krow_add('rev', 'Revenue ($)', lambda c, j: f"='P&L'!{c}{PREF['rev']}", USD)
krow_add('rev_town', 'Revenue per average live town ($)',
         lambda c, j: f"=IF(Rollout!{c}8=0,0,'P&L'!{c}{PREF['rev']}/Rollout!{c}8)", USD)
krow_add('gm', 'Gross margin', lambda c, j: f"='P&L'!{c}{PREF['gm']}", PCT)
krow_add('ebitda', 'EBITDA ($)', lambda c, j: f"='P&L'!{c}{PREF['ebitda']}", USD)
krow_add('ebitda_m', 'EBITDA margin', lambda c, j: f"='P&L'!{c}{PREF['ebitda_m']}", PCT)
krow_add('cash', 'Closing cash ($)', lambda c, j: f"='Cash Flow'!{c}13", USD)
krow_add('low_cash', 'Lowest month-end cash in year ($)',
         lambda c, j: f'=MIN(Monthly!{MC["close"]}{fy_rows(j)[0]}:{MC["close"]}{fy_rows(j)[1]})', USD)
krow_add('paid_fte', 'Paid staff at year end (FTE)', lambda c, j: f'=Headcount!{c}{HREF["paid_fte"]}', NUM1,
         'Founders are unpaid and not counted.')
krow_add('ai_rev', 'AI cost as % of revenue',
         lambda c, j: f"=IF('P&L'!{c}{PREF['rev']}=0,0,Costs!{c}{CREF['ai']}/'P&L'!{c}{PREF['rev']})", PCT,
         'CLAUDE.md §8 asks for an alert when AI cost per article jumps; this tracks the overall share.')
krow_add('central', 'Central costs (engineering, leadership, G&A, brand, fixed legal, tools) ($)',
         lambda c, j: f"=Costs!{c}{CREF['p_eng']}+Costs!{c}{CREF['p_lead']}+Costs!{c}{CREF['p_ga']}+Costs!{c}{CREF['brand']}+Assumptions!{c}${REF['legal_fixed']}+Costs!{c}{CREF['tools']}",
         USD, 'Bootstrapped: the people here are paid only once revenue covers them.')
krow_add('runrate', 'Run-rate EBITDA if every live town were mature ($)',
         lambda c, j: f"=Rollout!{c}7/'Unit Economics'!$G${UREF['surv']}*'Unit Economics'!$G${UREF['contrib']}-{c}{KREF['central']}",
         USD, 'Live towns x Age 5 contribution per surviving town, less central costs: the profit the '
         'network already built would earn once its young towns mature, with no new launches.')
krow_add('be_towns', 'Mature towns needed to cover central costs',
         lambda c, j: f"=IF('Unit Economics'!$G${UREF['contrib']}<=0,0,{c}{KREF['central']}/('Unit Economics'!$G${UREF['contrib']}/'Unit Economics'!$G${UREF['surv']}))",
         NUM)
krow += 1
section(K, krow, 'Single-town economics (from Unit Economics)')
krow += 1
K[f'A{krow}'] = 'Mature town revenue (Age 5, $/year)'
K[f'C{krow}'] = f"='Unit Economics'!G{UREF['rev']}"
K[f'C{krow}'].number_format = USD
K[f'C{krow}'].font = F_LINK
krow += 1
K[f'A{krow}'] = 'Mature town contribution (Age 5, $/year)'
K[f'C{krow}'] = f"='Unit Economics'!G{UREF['contrib']}"
K[f'C{krow}'].number_format = USD
K[f'C{krow}'].font = F_LINK
krow += 1
K[f'A{krow}'] = 'Mature contribution margin'
K[f'C{krow}'] = f"='Unit Economics'!G{UREF['contrib_m']}"
K[f'C{krow}'].number_format = PCT
K[f'C{krow}'].font = F_LINK
krow += 1
K[f'A{krow}'] = 'One-time launch cost per town ($)'
K[f'C{krow}'] = f"='Unit Economics'!C{UREF['launch_cost']}"
K[f'C{krow}'].number_format = USD
K[f'C{krow}'].font = F_LINK
krow += 1
K[f'A{krow}'] = 'Town payback (age in years when cumulative contribution turns positive)'
cum = UREF['cum']
K[f'C{krow}'] = (f"=IFERROR(MATCH(TRUE,INDEX('Unit Economics'!C{cum}:G{cum}>0,0),0),\"After Age 5\")")
K[f'C{krow}'].number_format = '0'
K[f'H{krow}'] = '1 = paid back within the launch year.'
K[f'H{krow}'].font = F_NOTE
krow += 1
K[f'A{krow}'] = 'First fiscal year with positive EBITDA'
K[f'C{krow}'] = (f"=IFERROR(INDEX(C4:G4,MATCH(TRUE,INDEX(C{KREF['ebitda']}:G{KREF['ebitda']}>0,0),0)),"
                 f"\"Not by FY2031\")")
krow += 1
K[f'A{krow}'] = 'Equity needed (deepest cumulative free cash flow)'
K[f'C{krow}'] = "='Cash Flow'!C17"
K[f'C{krow}'].number_format = USD
K[f'C{krow}'].font = F_LINK

# Cover
CV = wb.create_sheet('Cover', 0)
title(CV, 'TownCore long-range plan, FY2027-FY2031',
      'AI-written local news sites with a daily newsletter and a local business directory.')
setw(CV, {'A': 52, 'B': 4, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 30})
CV['A4'] = 'Scenario in use'
CV['C4'] = f'=INDEX({{"Downside","Base","Upside"}},1,{REF["scenario"]})'
CV['C4'].font = F_BOLD
header(CV, 6, 'Summary', YEARS)
summary = [
    ('Live towns at year end', 'Rollout!{c}7', NUM),
    ('Newsletter subscribers', f'KPIs!{{c}}{KREF["subs"]}', NUM),
    ('Revenue ($)', f"'P&L'!{{c}}{PREF['rev']}", USD),
    ('Gross margin', f"'P&L'!{{c}}{PREF['gm']}", PCT),
    ('EBITDA ($)', f"'P&L'!{{c}}{PREF['ebitda']}", USD),
    ('EBITDA margin', f"'P&L'!{{c}}{PREF['ebitda_m']}", PCT),
    ('Net income ($)', f"'P&L'!{{c}}{PREF['ni']}", USD),
    ('Closing cash ($)', "'Cash Flow'!{c}13", USD),
    ('Lowest month-end cash ($)', f'KPIs!{{c}}{KREF["low_cash"]}', USD),
    ('Paid staff at year end (FTE; founders unpaid)', f'KPIs!{{c}}{KREF["paid_fte"]}', NUM1),
    ('Run-rate EBITDA at maturity ($)', f'KPIs!{{c}}{KREF["runrate"]}', USD),
]
for k, (label, ref, fmt) in enumerate(summary):
    r = 7 + k
    CV[f'A{r}'] = label
    for c in COLS:
        CV[f'{c}{r}'] = '=' + ref.format(c=c)
        CV[f'{c}{r}'].number_format = fmt
        CV[f'{c}{r}'].font = F_LINK
r = 7 + len(summary) + 1
CV[f'A{r}'] = 'How to use'
CV[f'A{r}'].font = F_BOLD
notes = [
    'Change inputs only on the Assumptions tab: blue text is an input, yellow fill a key assumption.',
    'Switch scenarios with Assumptions "Active scenario" (1 Downside, 2 Base, 3 Upside).',
    'Black text is a formula; green text links to another tab. Don\'t overwrite formulas.',
    'Bootstrapped (owner, 2026-10-17): about $5k of opening cash, no outside money, founders unpaid. The '
    'Monthly sheet launches each month only what last month\'s cash allows (keeping the cash floor and a '
    'buffer) and only while no needed work is unfunded; staff are paid out of revenue.',
    'Unit Economics shows one average town by age (Age 1 = launch year). The Monthly sheet multiplies it '
    'by the towns launched in each month; Revenue, Costs and Cash Flow sum the months.',
    'Prices, story budgets, AI caps and directory counts come from the TownCore repo (CLAUDE.md and town '
    'config); everything marked "Estimate" in Assumptions column H is a planning assumption to validate.',
    'Not modeled: programmatic consent costs, a licensed sports-data feed (not yet bought), office space, '
    'debt, and acquisitions.',
]
for k, n in enumerate(notes):
    CV[f'A{r+1+k}'] = f'{k+1}. {n}'
    CV[f'A{r+1+k}'].font = F_BASE
r = r + len(notes) + 2
CV[f'A{r}'] = 'Color legend'
CV[f'A{r}'].font = F_BOLD
for k, (txt, font, fill) in enumerate([
    ('Input', F_INPUT, None), ('Key assumption', F_INPUT, FILL_KEY),
    ('Formula', F_BASE, None), ('Link to another tab', F_LINK, None)]):
    cell = CV[f'A{r+1+k}']
    cell.value = txt
    cell.font = font
    if fill:
        cell.fill = fill

# Fonts everywhere default to Arial; freeze panes and print setup
for ws in wb.worksheets:
    for rw in ws.iter_rows():
        for cell in rw:
            if cell.font is None or cell.font.name != ARIAL:
                f = cell.font
                cell.font = Font(name=ARIAL, size=f.size or 10, bold=f.bold, italic=f.italic, color=f.color)
    ws.sheet_view.showGridLines = False
    if ws.title not in ('Cover',):
        ws.freeze_panes = 'C5' if ws.title != 'Assumptions' else 'C4'

wb.save(OUT)
print('saved', OUT)
